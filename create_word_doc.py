from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = r"C:\Users\Administrator\Documents\默认项目\运动服视觉策略Skill转化路径.docx"

BLUE = RGBColor(46, 116, 181)
DARK = RGBColor(31, 77, 120)

def set_font(run, size=None, bold=None, color=None):
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    run._element.rPr.rFonts.set(qn('w:ascii'), 'Calibri')
    run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Calibri')
    if size: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if color: run.font.color.rgb = color

def format_paragraph(p, before=0, after=6, line=1.25):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line

def text_para(doc, text, before=0, after=6):
    p = doc.add_paragraph()
    format_paragraph(p, before, after)
    set_font(p.add_run(text), 11)
    return p

def heading(doc, text, level):
    p = doc.add_paragraph()
    format_paragraph(p, 18 if level == 1 else 14, 10 if level == 1 else 7)
    run = p.add_run(text)
    set_font(run, 16 if level == 1 else 13, True, BLUE if level == 1 else DARK)
    return p

def bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    format_paragraph(p, 0, 4)
    p.paragraph_format.left_indent = Inches(0.375)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    set_font(p.add_run(text), 11)
    return p

def num(doc, text):
    p = doc.add_paragraph(style='List Number')
    format_paragraph(p, 0, 4)
    p.paragraph_format.left_indent = Inches(0.375)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    set_font(p.add_run(text), 11)
    return p

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(1)
sec.bottom_margin = Inches(1)
sec.left_margin = Inches(1)
sec.right_margin = Inches(1)
sec.header_distance = Inches(0.492)
sec.footer_distance = Inches(0.492)

normal = doc.styles['Normal']
normal.font.name = 'Microsoft YaHei'
normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
normal.font.size = Pt(11)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
format_paragraph(title, 0, 3)
set_font(title.add_run('从设计经验到视觉策略 Skill'), 22, True, DARK)

sub = doc.add_paragraph()
format_paragraph(sub, 0, 12)
set_font(sub.add_run('以运动服产品为例的智能体能力沉淀路径'), 11, False, RGBColor(89, 89, 89))

heading(doc, '一、转化逻辑', 1)
text_para(doc, '将设计师的隐性经验拆解为可解释、可复用、可被智能体调用的规则，最终沉淀为面向具体品类的视觉策略 Skill。')

num(doc, '设计师经验：识别产品特征、用户需求与市场语境中的关键设计线索。')
num(doc, '设计判断路径：明确“为什么这样选择”的判断依据与优先级。')
num(doc, 'Skill 规则：将判断路径结构化为输入条件、策略规则与输出要求。')
num(doc, '智能体调用：根据产品输入匹配规则，生成可执行的视觉策略建议。')

heading(doc, '二、案例：运动服产品视觉策略 Skill', 1)

heading(doc, '1. 产品输入', 2)
text_para(doc, '智能体首先接收并理解以下产品信息：')
for item in ['品类：运动服', '面料：功能性、透气或弹力面料', '用户群体：关注运动表现与日常穿着体验的人群', '竞品定位：同类运动品牌的风格、价格带与传播方向']:
    bullet(doc, item)

heading(doc, '2. AI 学习的设计判断', 2)
text_para(doc, 'AI 不仅记录视觉结果，还需要学习设计选择背后的原因：')
for item in ['为什么选择户外场景：强化运动服的功能属性、环境适应性与真实使用感。', '为什么选择动态姿态：直观传递面料弹性、身体动作与产品性能。', '为什么选择特定色彩体系：与目标用户审美、品牌定位及竞品区隔保持一致。']:
    bullet(doc, item)

heading(doc, '3. 规则沉淀', 2)
text_para(doc, '把“产品输入”与“设计判断”转化为可调用规则：')
for item in ['当品类为运动服，优先呈现与运动功能相关的真实使用场景。', '当面料强调弹力、透气或支撑性，优先采用可展现身体动作的动态姿态。', '根据用户群体与竞品定位，确定具有辨识度且符合运动语境的色彩体系。']:
    bullet(doc, item)

heading(doc, '三、最终输出', 1)
text_para(doc, '最终形成“运动服视觉策略 Skill”：')
for item in ['输入：品类、面料、用户群体、竞品定位。', '推理：匹配场景、姿态、色彩等视觉判断规则。', '输出：可直接用于视觉方案与 AI 生产的策略建议。']:
    bullet(doc, item)

note = doc.add_paragraph()
format_paragraph(note, 10, 0)
r = note.add_run('核心价值：把“设计师知道怎么做”转化为“智能体知道何时、为何、如何做”。')
set_font(r, 11, True, DARK)

doc.save(OUT)
print(OUT)
